from __future__ import annotations

from pathlib import Path
from tempfile import TemporaryDirectory

from docx import Document
from docx.shared import Pt

from services.document_model import build_document_model
from services.docx_analyzer import analyze_docx
from services.executor_adapter import execute_plan
from services.planner import ExecutionPlan, PlanStep
from services.rule_engine import Rule, RuleSource
from services.verifier import verify_output


def check(name: str, condition: bool) -> None:
    if not condition:
        raise AssertionError(f"{name} FAIL")
    print(f"{name} PASS")


def make_source(path: Path) -> None:
    document = Document()
    document.add_paragraph("格式验证示例论文")
    document.add_paragraph("摘要：用于验证低风险格式修改与 provenance 记录。")
    document.add_paragraph("关键词：格式；验证；论文")
    document.add_paragraph("1 引言")
    for text in ["第一段正文用于验证行距。", "第二段正文用于验证行距。", "第三段正文用于验证行距。"]:
        paragraph = document.add_paragraph(text)
        paragraph.paragraph_format.line_spacing = 1.0
        for run in paragraph.runs:
            run.font.size = Pt(9)
    document.add_paragraph("参考文献")
    document.add_paragraph("[1] 示例. 格式验证[J]. 测试, 2024.")
    document.save(path)


def main() -> None:
    with TemporaryDirectory() as directory:
        source, output = Path(directory) / "source.docx", Path(directory) / "output.docx"
        make_source(source)
        model = build_document_model(source, classification={"document_type": "academic_paper"})
        before = analyze_docx(source)
        rule = Rule("body-line_spacing", "body", "line_spacing", 1.5, RuleSource.DEFAULT, "test", 0.9, "low", True)
        safe_step = PlanStep("step-body-spacing", "apply_body_paragraph_format", "body", rule.id, rule.evidence, "low", True, [], "planned", "test", target_locator={"kind": "paragraph_indices", "indices": model.body_paragraph_indices, "semantic_role": "body"})
        risky_step = PlanStep("review-references", "human_review", "references", "reference-risk", "test", "high_risk", False, [], "human_review_required", "test")
        plan = ExecutionPlan("plan-provenance", model.document_id, [safe_step, risky_step], [], True)
        execution = execute_plan(plan, source, output, None, [rule])

        check("low_risk_step_executed", safe_step.id in execution.executed_step_ids)
        check("high_risk_step_not_faked", risky_step.id in execution.unsupported_step_ids and execution.step_statuses[risky_step.id] == "unsupported")
        check("provenance_before_after", bool(execution.changes) and all(change["before"] == 1.0 and change["after"] == 1.5 for change in execution.changes))
        check("provenance_plan_rule_link", all(change["plan_step_id"] == safe_step.id and change["rule_id"] == rule.id and change["target"].get("paragraph_index") is not None for change in execution.changes))
        check("output_created", output.exists() and output.stat().st_size > 0)

        verification = verify_output(model, before, output, None, [rule], plan, execution.unsupported_step_ids, execution.changes)
        check("verifier_provenance_enriched", bool(verification.provenance_changes) and all(change["verification_scope"] == "target" and change["verification_status"] == "verified" and "verification_evidence" in change for change in verification.provenance_changes))
        check("structure_integrity_guard", verification.structural_integrity["status"] in {"SAFE", "WARNING"})
    print("P1_EXECUTOR_PROVENANCE PASS")


if __name__ == "__main__":
    main()
