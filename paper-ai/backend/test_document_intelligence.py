from __future__ import annotations

import json
from pathlib import Path

from docx import Document

from services.agent_pipeline import run_agent_pipeline
from services.document_intelligence import analyze_document


def _make_doc(path: Path, *, language: str) -> None:
    document = Document()
    if language == "zh":
        values = ["基于知识图谱的学术文档分析研究", "摘要", "本文提出一种用于论文结构分析的方法。", "关键词：知识图谱；文档分析", "第一章 绪论", "研究背景与问题定义。", "第二章 相关研究", "介绍已有研究工作。", "第三章 实验方法", "本文设计实验并报告结果。", "第四章 实验结果", "结果表明方法有效。", "第五章 结语", "总结本文工作并讨论未来方向。", "参考文献", "[1] 作者. 文献标题. 2025."]
    else:
        values = ["Document Intelligence for Academic Papers", "Abstract", "This paper presents a method for document structure analysis.", "Keywords: document analysis; academic paper", "1 Introduction", "We define the research problem and background.", "2 Related Work", "Prior studies are reviewed in this section.", "3 Methodology", "We design an experiment to evaluate the method.", "4 Results", "The results demonstrate that the method is effective.", "5 Conclusion", "We summarize the work and discuss future work.", "References", "[1] Author. Paper title. 2025."]
    heading_values = {"第一章 绪论", "第二章 相关研究", "第三章 实验方法", "第四章 实验结果", "第五章 结语", "1 Introduction", "2 Related Work", "3 Methodology", "4 Results", "5 Conclusion", "参考文献", "References"}
    for value in values:
        paragraph = document.add_paragraph(value)
        if value in heading_values:
            paragraph.style = "Heading 1"
    document.add_table(rows=1, cols=2).rows[0].cells[0].text = "Table 1"
    document.save(path)


def test_chinese_sections_and_semantics(tmp_path: Path) -> None:
    path = tmp_path / "zh.docx"
    _make_doc(path, language="zh")
    analysis = analyze_document(path)
    types = {section.type for section in analysis.sections}
    assert {"title", "abstract", "keywords", "introduction", "related_work", "methodology", "experiment", "conclusion", "references"}.issubset(types)
    assert any(item.semantic_type == "reference" for item in analysis.elements)
    assert all(0 <= item.confidence <= 1 for item in analysis.elements)


def test_english_sections_and_confidence(tmp_path: Path) -> None:
    path = tmp_path / "en.docx"
    _make_doc(path, language="en")
    analysis = analyze_document(path)
    assert analysis.metadata["language"] == "en"
    assert {"introduction", "related_work", "methodology", "experiment", "conclusion", "references"}.issubset({item.type for item in analysis.sections})
    assert 0 < analysis.confidence <= 1
    assert all(0 <= item.confidence <= 1 for item in analysis.sections)
    assert analysis.metadata["table_count"] == 1


def test_unknown_heading_has_lower_confidence(tmp_path: Path) -> None:
    path = tmp_path / "unknown.docx"
    document = Document()
    document.add_paragraph("A Study")
    document.add_paragraph("7 Unnamed Part", style="Heading 1")
    document.add_paragraph("Some body text that cannot determine the section role.")
    document.save(path)
    analysis = analyze_document(path)
    unknown = [item for item in analysis.sections if item.title == "7 Unnamed Part"]
    assert unknown and unknown[0].type == "unknown"
    assert unknown[0].confidence < 0.8


def test_legacy_pipeline_and_task_state_compatibility(tmp_path: Path) -> None:
    path = tmp_path / "paper.docx"
    _make_doc(path, language="en")
    output_dir = tmp_path / "outputs"
    output_dir.mkdir()
    result = run_agent_pipeline(path, output_dir, mode="local")
    assert result["status"] == "ok"
    assert result["document_analysis"]["sections"]
    state = json.loads(Path(result["task_state_path"]).read_text(encoding="utf-8"))
    assert state["document_analysis"]["document_type"] in {"academic_paper", "unknown"}
    assert state["ai_score"] is None
    assert state["ai_used"] is False
