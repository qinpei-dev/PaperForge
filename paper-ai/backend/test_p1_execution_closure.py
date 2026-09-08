from __future__ import annotations

import tempfile
from pathlib import Path

from docx import Document

from services.executor_adapter import execute_plan
from services.governance import decide, human_review_from_decision
from services.planner import ExecutionPlan, PlanStep, normalize_execution_plan
from services.rule_engine import Rule, RuleSource
from services.verifier import verify_output
from services.document_model import build_document_model
from services.docx_analyzer import analyze_docx


def rule(rule_id: str, prop: str, expected):
    return Rule(rule_id, "body", prop, expected, RuleSource.DEFAULT, "test", 1.0, "low", True)


def step(step_id: str, rule_id: str, prop: str, expected, locator=None):
    return PlanStep(step_id, "apply_body_font" if prop in {"font_name", "font_size"} else "apply_body_paragraph_format", "body", rule_id, "test", "low", True, [], "planned", "test", target_locator=locator or {"kind": "paragraph_indices", "indices": [4], "semantic_role": "body", "target_type": "body_paragraph"}, expected=expected)


def make_doc(path: Path):
    doc = Document()
    for index in range(6):
        doc.add_paragraph(f"paragraph {index}")
    doc.save(path)


def check(name, condition):
    print(f"{name} {'PASS' if condition else 'FAIL'}")
    if not condition:
        raise AssertionError(name)


def main():
    with tempfile.TemporaryDirectory() as tmp:
        source, output = Path(tmp) / "source.docx", Path(tmp) / "output.docx"
        make_doc(source)
        rules = [rule("body-font_size", "font_size", 12), rule("body-alignment", "alignment", 3), rule("body-first_line_indent_cm", "first_line_indent_cm", 0.74)]
        plan = ExecutionPlan("plan-test", "doc-test", [step("font", rules[0].id, "font_size", 12), step("align", rules[1].id, "alignment", 3), step("indent", rules[2].id, "first_line_indent_cm", 0.74)], [], False)
        normalized, summary = normalize_execution_plan(plan)
        check("multi_field_kept", len(normalized.steps) == 3 and summary["conflicts"] == 0)
        execution = execute_plan(normalized, source, output, None, rules)
        verification = verify_output(build_document_model(source, classification={"document_type": "academic_paper"}), analyze_docx(source), output, None, rules, normalized, execution.unsupported_step_ids, execution.changes)
        check("multi_field_verified", verification.verification_summary["verified"] == 3 and verification.verification_summary["failed"] == 0)
        check("non_target_protected", Document(output).paragraphs[3].text == "paragraph 3" and Document(output).paragraphs[5].text == "paragraph 5")

        duplicate_plan = ExecutionPlan("plan-dup", "doc-test", [step("same-a", rules[0].id, "font_size", 12), step("same-b", rules[0].id, "font_size", 12)], [], False)
        _, duplicate_summary = normalize_execution_plan(duplicate_plan)
        check("duplicate_deduplicated", duplicate_summary["duplicates"] == 1)

        conflict_plan = ExecutionPlan("plan-conflict", "doc-test", [step("center", "body-alignment", "alignment", 1), step("justify", "body-alignment", "alignment", 3)], [], False)
        conflict_normalized, conflict_summary = normalize_execution_plan(conflict_plan)
        conflict_execution = execute_plan(conflict_normalized, source, Path(tmp) / "conflict.docx", None, [rule("body-alignment", "alignment", 1)])
        check("conflict_detected", conflict_summary["conflicts"] == 1 and conflict_execution.conflict_step_ids)
        check("conflict_not_silent", all(item["verification_status"] == "unsupported" for item in conflict_execution.changes))
        conflict_decision = decide({"verification_summary": {"total": 2, "verified": 0, "failed": 0, "unsupported": 2}, "conflict_summary": conflict_summary, "passed": True}, replan_count=0, max_replans=1, plan_requires_review=False)
        check("conflict_requires_hitl", conflict_decision["action"] == "HUMAN_REVIEW")

        failure_decision = decide({"passed": True, "verification_summary": {"total": 1, "verified": 0, "failed": 1, "unsupported": 0}, "structural_integrity": {"status": "SAFE"}}, replan_count=0, max_replans=1, plan_requires_review=False)
        check("target_failure_overrides_analyzer", failure_decision["action"] == "HUMAN_REVIEW")

        unsupported_plan = ExecutionPlan("plan-unsupported", "doc-test", [step("missing", "body-font_size", "font_size", 12, {"kind": "semantic_role", "role": "body"})], [], False)
        unsupported_execution = execute_plan(unsupported_plan, source, Path(tmp) / "unsupported.docx", None, rules[:1])
        check("unsupported_continues", Path(unsupported_execution.output_path).exists() and unsupported_execution.unsupported_step_ids)
        unsupported_decision = decide({"passed": False, "verification_summary": {"total": 1, "verified": 0, "failed": 0, "unsupported": 1}, "unsupported_step_ids": ["missing"], "provenance_changes": unsupported_execution.changes, "structural_integrity": {"status": "SAFE"}}, replan_count=0, max_replans=1, plan_requires_review=False)
        review = human_review_from_decision(unsupported_decision, unsupported_plan)
        check("unsupported_hitl_target", unsupported_decision["action"] == "HUMAN_REVIEW" and review.affected_targets)

    print("P1_EXECUTION_CLOSURE PASS")


if __name__ == "__main__":
    main()
