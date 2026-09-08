from pathlib import Path
from tempfile import TemporaryDirectory

from docx import Document

from services.content_review import apply_content_review, apply_content_suggestion, review_content
from services.review_evidence import aggregate_review_evidence, content_score_summary


def make_doc(path: Path, text: str) -> None:
    doc = Document()
    doc.add_paragraph(text)
    doc.save(path)


def main() -> None:
    with TemporaryDirectory() as tmp:
        root = Path(tmp)
        source = root / "source.docx"
        make_doc(source, "很多的表达需要调整。")
        review = review_content(source)
        suggestion = next(item for item in review["issues"] if item["action_policy"] == "SUGGEST_ONLY")
        assert suggestion["issue_id"] and suggestion["locator"]["paragraph_index"] == 0

        untouched = root / "untouched.docx"
        applied_review = apply_content_review(source, untouched, review)
        assert Document(untouched).paragraphs[0].text == "很多的表达需要调整。"
        assert applied_review["provenance"]["suggestions"][0]["verification_status"] == "original_unchanged"
        score = content_score_summary(applied_review)
        assert score["verified_changes_count"] == 0 and score["delta"] == 0

        auto_source = root / "auto.docx"
        make_doc(auto_source, "这是  一个测试。。")
        auto_review = apply_content_review(auto_source, root / "auto-output.docx", review_content(auto_source))
        auto_score = content_score_summary(auto_review)
        assert auto_score["verified_changes_count"] == 1 and auto_score["delta"] > 0

        hitl_source = root / "hitl.docx"
        make_doc(hitl_source, "实验结果为 95%，结论需要人工确认。")
        hitl_review = apply_content_review(hitl_source, root / "hitl-output.docx", review_content(hitl_source))
        hitl_score = content_score_summary(hitl_review)
        assert hitl_score["verified_changes_count"] == 0 and hitl_score["delta"] == 0

        confirmed = root / "confirmed.docx"
        accepted = apply_content_suggestion(source, confirmed, suggestion)
        assert accepted["status"] == "ok", accepted
        assert accepted["change"]["accepted_by_user"] is True
        assert accepted["change"]["verification_status"] == "verified"
        assert Document(confirmed).paragraphs[0].text == suggestion["suggested_text"]
        evidence = aggregate_review_evidence(content_review={**applied_review, "provenance": {**applied_review["provenance"], "accepted": [accepted["change"]]}})
        assert evidence["review_summary"]["content"]["accepted"] == 1
        assert evidence["review_summary"]["overall"]["total_verified_changes"] == 1

        stale_source = root / "stale.docx"
        make_doc(stale_source, "已经被其他修改改变。")
        stale = apply_content_suggestion(stale_source, root / "stale-output.docx", suggestion)
        assert stale["status"] == "conflict", stale

    print("P2_CLOSURE PASS")


if __name__ == "__main__":
    main()
