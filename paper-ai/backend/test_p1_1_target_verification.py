from __future__ import annotations

from pathlib import Path
from tempfile import TemporaryDirectory

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from services.document_model import build_document_model
from services.docx_analyzer import analyze_docx
from services.executor_adapter import execute_plan
from services.planner import ExecutionPlan, PlanStep, reliable_caption_locator, reliable_heading_locator
from services.rule_engine import Rule, RuleSource
from services.verifier import verify_output


def check(name: str, condition: bool) -> None:
    if not condition:
        raise AssertionError(f"{name} FAIL")
    print(f"{name} PASS")


def make_source(path: Path) -> None:
    document = Document()
    document.add_paragraph("目标级验证示例")
    heading = document.add_paragraph("1 引言")
    heading.style = document.styles["Heading 1"]
    heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
    heading.runs[0].font.size = Pt(10)
    document.add_paragraph("正文保留原始内容，见图 1 和表 1。")
    figure = document.add_paragraph("图 1 系统结构")
    figure.runs[0].font.size = Pt(8)
    table = document.add_paragraph("表 1 样本统计")
    table.runs[0].font.size = Pt(8)
    document.add_paragraph("参考文献")
    document.add_paragraph("[1] 示例文献[J]. 测试, 2026.")
    document.save(path)


def main() -> None:
    with TemporaryDirectory() as directory:
        source, output = Path(directory) / "source.docx", Path(directory) / "output.docx"
        make_source(source)
        model = build_document_model(source, classification={"document_type": "academic_paper"})
        heading_locator = reliable_heading_locator(model, "heading:Heading 1")
        figure_locator, table_locator = reliable_caption_locator(model, "figure"), reliable_caption_locator(model, "table")
        check("reliable_heading_locator", heading_locator and heading_locator["indices"] == [1] and heading_locator["confidence"] >= 0.9)
        check("figure_caption_locator", figure_locator and figure_locator["indices"] == [3] and figure_locator["semantic_role"] == "figure_caption")
        check("table_caption_locator", table_locator and table_locator["indices"] == [4] and table_locator["semantic_role"] == "table_caption")

        heading_rule = Rule("heading-size", "heading:Heading 1", "font_size", 16, RuleSource.TEMPLATE, "test", .9, "low", True)
        figure_rule = Rule("figure-size", "caption:figure", "font_size", 10.5, RuleSource.DEFAULT, "test", .72, "low", True)
        table_rule = Rule("table-align", "caption:table", "alignment", WD_ALIGN_PARAGRAPH.CENTER, RuleSource.DEFAULT, "test", .72, "low", True)
        plan = ExecutionPlan("plan-p1-1", model.document_id, [
            PlanStep("heading", "apply_heading_format", "heading", heading_rule.id, "test", "low", True, [], "planned", "test", target_locator=heading_locator),
            PlanStep("figure", "apply_caption_format", "caption:figure", figure_rule.id, "test", "low", True, [], "planned", "test", target_locator=figure_locator),
            PlanStep("table", "apply_caption_format", "caption:table", table_rule.id, "test", "low", True, [], "planned", "test", target_locator=table_locator),
            PlanStep("bad", "apply_heading_format", "heading", heading_rule.id, "test", "low", True, [], "planned", "test", target_locator={"kind": "paragraph_indices", "indices": [99], "semantic_role": "heading"}),
            PlanStep("missing", "apply_caption_format", "caption:figure", figure_rule.id, "test", "low", True, [], "planned", "test", target_locator={}),
        ], [], False)
        execution = execute_plan(plan, source, output, None, [heading_rule, figure_rule, table_rule])
        check("target_steps_only", set(execution.executed_step_ids) == {"heading", "figure", "table"})
        check("missing_or_out_of_range_not_faked", execution.step_statuses["bad"] == "skipped" and execution.step_statuses["missing"] == "unsupported")
        result = Document(output)
        check("caption_and_heading_only_targeted", result.paragraphs[2].text == "正文保留原始内容，见图 1 和表 1。" and result.paragraphs[3].text == "图 1 系统结构" and result.paragraphs[4].text == "表 1 样本统计")
        check("before_after_provenance", all(change["before"] != change["after"] and change["target"]["paragraph_index"] in {1, 3, 4} for change in execution.changes))
        verification = verify_output(model, analyze_docx(source), output, None, [heading_rule, figure_rule, table_rule], plan, execution.unsupported_step_ids, execution.changes)
        check("target_verification", all(change["verification_scope"] == "target" and change["verification_status"] == "verified" and change["verification_evidence"]["source"] == "re-read output DOCX" for change in verification.provenance_changes))
        low_confidence = type(model.headings[0])(level=1, text="1 模糊标题", paragraph_index=1, style=None, confidence=.72)
        altered = type(model)(**{**model.to_dict(), "headings": [low_confidence]})
        check("low_confidence_heading_no_locator", reliable_heading_locator(altered, "heading:Heading 1") is None)
        check("structure_intact", verification.structural_integrity["status"] in {"SAFE", "WARNING"})
    print("P1_1_TARGET_VERIFICATION PASS")


if __name__ == "__main__":
    main()
