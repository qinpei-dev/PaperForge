from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from services.agent_pipeline import run_agent_pipeline
from services.ai_reasoning import generate_reasoning
from services.template_intelligence import analyze_template


def _make_template(path: Path) -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1.25)
    section.right_margin = Inches(1.25)
    header = section.header.paragraphs[0]
    header.text = "PaperForge University Thesis"
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    for text in ("毕业论文", "姓名：________", "学号：________", "目录", "摘要", "正文", "参考文献", "附录"):
        paragraph = document.add_paragraph(text)
        paragraph.style = "Heading 1" if text in {"目录", "摘要", "正文", "参考文献", "附录"} else "Normal"
    document.add_table(rows=2, cols=2).cell(0, 0).text = "姓名"
    document.save(path)


def test_template_sections_rules_and_protected_regions(tmp_path: Path) -> None:
    path = tmp_path / "template.docx"
    _make_template(path)
    analysis = analyze_template(path)
    assert {item["type"] for item in analysis.sections} >= {"cover", "abstract", "table_of_contents", "body", "references", "appendix"}
    properties = {item["property"] for item in analysis.rules}
    assert {"font", "size", "bold", "alignment", "spacing", "margins"} <= properties
    protected_types = {item["type"] for item in analysis.protected_regions}
    assert {"cover_fields", "fixed_headers", "fixed_tables"} <= protected_types
    assert 0 < analysis.confidence <= 1


def test_template_analysis_is_consumed_by_reasoning_without_breaking_old_call() -> None:
    old = generate_reasoning({"confidence": 0.8}, [{"issue_type": "heading_mismatch", "score": 80}], [])
    new = generate_reasoning(
        {"confidence": 0.8},
        [{"issue_type": "heading_mismatch", "score": 80}],
        [],
        {"confidence": 0.9, "sections": [{"type": "body"}], "protected_regions": [{"type": "fixed_headers"}]},
    )
    assert old and new
    assert new[0]["context"]["template_sections"] == ["body"]
    assert new[0]["context"]["protected_region_types"] == ["fixed_headers"]


def test_pipeline_and_task_state_backward_compatible_without_template(tmp_path: Path) -> None:
    paper = tmp_path / "paper.docx"
    document = Document()
    document.add_paragraph("A short academic paper title")
    document.add_paragraph("Abstract")
    document.add_paragraph("This paper presents a method for testing template intelligence.")
    document.add_paragraph("References", style="Heading 1")
    document.add_paragraph("[1] Author. Title. 2026.")
    document.save(paper)
    output_dir = tmp_path / "outputs"
    output_dir.mkdir()
    result = run_agent_pipeline(paper, output_dir, mode="local")
    assert result["status"] == "ok"
    assert result["template_analysis"]["template_type"] == "none"
    state = json.loads(Path(result["task_state_path"]).read_text(encoding="utf-8"))
    assert state["template_analysis"]["template_type"] == "none"
    assert state["ai_score"] is None
    assert state["ai_used"] is False

