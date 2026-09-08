from __future__ import annotations

from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any

from docx import Document

from .document_model import DocumentModel, build_document_model
from .docx_analyzer import analyze_docx
from .governance import compare_structure
from .planner import ExecutionPlan
from .rule_engine import Rule


@dataclass
class VerificationResult:
    passed: bool
    score_before: int
    score_after: int
    verified_rules: list[str]
    failed_rules: list[str]
    structural_integrity: dict[str, Any]
    new_issues: list[str]
    unresolved_issues: list[str]
    risk_level: str
    recommendation: str
    fixable: bool
    unsupported_step_ids: list[str]
    after_document_model: dict[str, Any]
    after_analysis: dict[str, Any]
    provenance_changes: list[dict[str, Any]]
    verification_summary: dict[str, int]

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)


def verify_output(before_model: DocumentModel, before_analysis: dict[str, Any], output: Path, template_path: Path | None, rules: list[Rule], plan: ExecutionPlan, unsupported_step_ids: list[str], changes: list[dict[str, Any]] | None = None) -> VerificationResult:
    after_model = build_document_model(output, classification={"document_type": before_model.document_type})
    after_analysis = analyze_docx(output, template_path=template_path)
    before_score = int(before_analysis["report"]["score"])
    after_score = int(after_analysis["report"]["score"])
    breakdown = {item["key"]: item for item in after_analysis["report"].get("breakdown", [])}
    verified, failed = [], []
    for rule in rules:
        key = "heading_format" if rule.target.startswith("heading:") else {"page": "page_margin", "body": "body_font", "references": "references"}.get(rule.target)
        if key and key in breakdown:
            (verified if int(breakdown[key].get("score", 0)) >= 90 else failed).append(rule.id)
    integrity = compare_structure(before_model.to_dict(), after_model.to_dict())
    risks = [item.get("message", "") for section in (after_analysis.get("reference_check") or {}, after_analysis.get("figure_table_check") or {}) for item in section.get("risk_items") or []]
    fixable = bool(failed) and all(rule.auto_fixable and rule.risk_level == "low" for rule in rules if rule.id in failed)
    passed = not failed and integrity["status"] != "UNSAFE" and not unsupported_step_ids
    risk = "high_risk" if integrity["status"] == "UNSAFE" or unsupported_step_ids else ("medium" if failed or risks else "low")
    provenance_changes = []
    output_document = Document(output)
    for change in changes or []:
        entry = dict(change)
        if entry.get("verification_scope") == "document":
            entry["verification_status"] = "verified" if integrity["status"] != "UNSAFE" else "verification_failed"
            entry["verification_evidence"] = {"structural_integrity": integrity["status"]}
        elif entry.get("verification_scope") == "target":
            target_index = (entry.get("target") or {}).get("paragraph_index")
            property_name = entry.get("rule_type")
            if entry.get("status") == "unsupported":
                entry["verification_status"] = "unsupported"
                entry["verification_evidence"] = entry.get("verification_evidence") or {"reason": "executor did not receive a reliable target"}
                provenance_changes.append(entry)
                continue
            if not isinstance(target_index, int) or not 0 <= target_index < len(output_document.paragraphs):
                entry["verification_status"] = "failed"
                entry["verification_evidence"] = {"reason": "target paragraph index is unavailable after re-reading output"}
            else:
                from .docx_formatter import describe_low_risk_paragraph_property
                actual = describe_low_risk_paragraph_property(output_document.paragraphs[target_index], property_name)
                expected = entry.get("expected", entry.get("after"))
                target_verified = values_match(actual, expected, property_name)
                entry["verification_status"] = "verified" if target_verified else "failed"
                entry["verification_evidence"] = {
                    "paragraph_index": target_index,
                    "expected": {property_name: expected},
                    "actual": {property_name: actual},
                    "source": "re-read output DOCX",
                }
        else:
            rule_verified = entry.get("rule_id") in verified
            entry["verification_status"] = "verified" if rule_verified else "verification_failed" if entry.get("rule_id") in failed else "unresolved"
            entry["verification_scope"] = "rule"
            entry["verification_evidence"] = {"rule_id": entry.get("rule_id"), "after_score": breakdown.get("body_font" if entry.get("rule_type") in {"font_name", "font_size"} else "spacing_indent", {}).get("score")}
        provenance_changes.append(entry)
    summary = {"total": 0, "verified": 0, "failed": 0, "unsupported": 0}
    for entry in provenance_changes:
        if entry.get("verification_scope") != "target":
            continue
        summary["total"] += 1
        status = entry.get("verification_status")
        if status == "verified": summary["verified"] += 1
        elif status in {"failed", "verification_failed"}: summary["failed"] += 1
        elif status in {"unsupported", "skipped"}: summary["unsupported"] += 1
    summary["unsupported"] = max(summary["unsupported"], len(unsupported_step_ids))
    if summary["failed"]:
        passed = False
    return VerificationResult(passed, before_score, after_score, verified, failed, integrity, risks, failed, risk, "通过" if passed else "需要决策引擎处理未满足规则或风险。", fixable, unsupported_step_ids, after_model.to_dict(), after_analysis, provenance_changes, summary)


def values_match(actual: Any, expected: Any, property_name: str) -> bool:
    if property_name == "alignment":
        # python-docx renders enum values as strings; expected values can come
        # from either a template enum or the serialized rule integer.
        if expected is None:
            return "JUSTIFY" in str(actual)
        return str(actual) == str(expected) or str(expected) in str(actual)
    if isinstance(actual, (int, float)) and isinstance(expected, (int, float)):
        return abs(float(actual) - float(expected)) <= 0.02
    return actual == expected
