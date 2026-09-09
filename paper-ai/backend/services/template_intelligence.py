from __future__ import annotations

import re
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any

from docx import Document

from .template_extractor import (
    extract_body_style,
    extract_heading_styles,
    extract_template_profile,
    point_value,
    safe_margin,
    safe_size,
    safe_spacing,
)


TEMPLATE_SECTION_TERMS: dict[str, tuple[str, ...]] = {
    "cover": ("封面", "cover", "课程论文", "毕业论文", "论文题目"),
    "abstract": ("摘要", "abstract", "summary"),
    "table_of_contents": ("目录", "contents", "table of contents"),
    "body": ("正文", "引言", "绪论", "introduction", "研究方法", "method", "实验", "results", "结论", "conclusion"),
    "references": ("参考文献", "references", "bibliography", "works cited"),
    "appendix": ("附录", "appendix", "annex"),
}

_FIELD_MARKERS = ("姓名", "学号", "学院", "专业", "班级", "课程", "指导教师", "student", "author", "school", "major")
_INSTRUCTION_MARKERS = ("填写", "请在", "模板", "format", "instruction", "font", "字号", "字体", "行距")


@dataclass(frozen=True)
class TemplateAnalysis:
    template_type: str
    sections: list[dict[str, Any]] = field(default_factory=list)
    rules: list[dict[str, Any]] = field(default_factory=list)
    protected_regions: list[dict[str, Any]] = field(default_factory=list)
    confidence: float = 0.0

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)


def analyze_template(path: Path | None) -> TemplateAnalysis:
    """Read-only, conservative understanding of an uploaded DOCX template."""
    if not path:
        return TemplateAnalysis("none")
    try:
        document = Document(path)
    except Exception:
        return TemplateAnalysis("unknown", confidence=0.0)

    profile = extract_template_profile(path) or {}
    sections = _identify_sections(document)
    rules = _extract_rules(document, profile)
    protected = _identify_protected_regions(document, sections)
    template_type = _classify_template_type(document, sections, protected)
    signals = [item["confidence"] for item in sections + rules + protected]
    confidence = _clamp(sum(signals) / len(signals) if signals else 0.35)
    return TemplateAnalysis(template_type, sections, rules, protected, confidence)


def _identify_sections(document: Any) -> list[dict[str, Any]]:
    result: list[dict[str, Any]] = []
    paragraphs = list(document.paragraphs)
    first_content = next((i for i, p in enumerate(paragraphs) if p.text.strip()), None)
    for index, paragraph in enumerate(paragraphs):
        text = paragraph.text.strip()
        if not text:
            continue
        kind, confidence = _section_kind(text, index, first_content, paragraph)
        if kind == "body" and not _is_heading_like(paragraph, text):
            continue
        result.append({"type": kind, "title": text, "paragraph_id": index, "confidence": confidence})
    if first_content is not None and not any(item["type"] == "cover" for item in result):
        first = paragraphs[first_content].text.strip()
        if _looks_like_cover(first, paragraphs[: first_content + 8]):
            result.insert(0, {"type": "cover", "title": first, "paragraph_id": first_content, "confidence": 0.72})
    return result


def _section_kind(text: str, index: int, first_content: int | None, paragraph: Any) -> tuple[str, float]:
    normalized = re.sub(r"^[第\d一二三四五六七八九十百]+\s*[章节、.：:]?\s*", "", text).strip().lower()
    matches = [
        (kind, term)
        for kind, terms in TEMPLATE_SECTION_TERMS.items()
        for term in terms
        if normalized == term.lower() or normalized.startswith(tuple(term.lower() + suffix for suffix in ("：", ":", " ")))
    ]
    if matches:
        kind, term = max(matches, key=lambda item: len(item[1]))
        return kind, 0.94 if _is_heading_like(paragraph, text) else 0.82
    if first_content is not None and index <= first_content + 8 and _looks_like_cover(text, [paragraph]):
        return "cover", 0.68
    return "body", 0.42


def _extract_rules(document: Any, profile: dict[str, Any]) -> list[dict[str, Any]]:
    rules: list[dict[str, Any]] = []

    def add(rule_id: str, target: str, property_name: str, value: Any, confidence: float, evidence: str) -> None:
        if value is not None or property_name == "alignment":
            rules.append({"id": rule_id, "target": target, "property": property_name, "value": value, "confidence": _clamp(confidence), "evidence": evidence})

    normal = profile.get("normal") or {}
    add("body-font", "body", "font", normal.get("font_name"), 0.90, "Normal style")
    add("body-size", "body", "size", normal.get("font_size"), 0.90, "Normal style")
    body = profile.get("body") or {}
    add("body-alignment", "body", "alignment", _alignment_name(body.get("alignment")), 0.82, "representative body paragraph")
    add("body-spacing", "body", "spacing", {key: body.get(key) for key in ("line_spacing", "space_before_pt", "space_after_pt", "first_line_indent_cm")}, 0.84, "representative body paragraph")
    heading = profile.get("heading") or {}
    for style_name, value in heading.items():
        target = f"heading:{style_name}"
        add(f"{style_name}-font", target, "font", value.get("font_name"), 0.88, style_name)
        add(f"{style_name}-size", target, "size", value.get("font_size"), 0.88, style_name)
        add(f"{style_name}-bold", target, "bold", value.get("bold"), 0.88, style_name)
    margins = profile.get("margins") or {}
    add("page-margins", "page", "margins", {side: margins.get(side) for side in ("top", "bottom", "left", "right")}, 0.92, "first document section")
    return rules


def _identify_protected_regions(document: Any, sections: list[dict[str, Any]]) -> list[dict[str, Any]]:
    protected: list[dict[str, Any]] = []
    paragraphs = list(document.paragraphs)
    cover_ids = [i for i, p in enumerate(paragraphs) if p.text.strip() and (_looks_like_cover(p.text.strip(), paragraphs[: min(len(paragraphs), 12)]) or _contains_field(p.text))]
    if cover_ids:
        protected.append({"type": "cover_fields", "paragraph_ids": sorted(set(cover_ids)), "confidence": 0.91, "evidence": "cover labels or fillable fields"})
    for section_index, section in enumerate(document.sections):
        header_ids = [p.text.strip() for p in section.header.paragraphs if p.text.strip()]
        if header_ids:
            protected.append({"type": "fixed_headers", "section_id": section_index, "content": header_ids, "confidence": 0.95, "evidence": "non-empty Word header"})
    for table_index, table in enumerate(document.tables):
        cell_text = " ".join(cell.text.strip() for row in table.rows for cell in row.cells if cell.text.strip())
        if cell_text and (_contains_field(cell_text) or len(table.rows) <= 4 and len(table.columns) <= 4):
            protected.append({"type": "fixed_tables", "table_id": table_index, "confidence": 0.76, "evidence": "small fixed-layout table or cover fields"})
    return protected


def _classify_template_type(document: Any, sections: list[dict[str, Any]], protected: list[dict[str, Any]]) -> str:
    text = " ".join(p.text.lower() for p in document.paragraphs if p.text.strip())
    if "期刊" in text or "journal" in text:
        return "journal_article"
    if "课程" in text or "course" in text or any(item["type"] == "cover" for item in sections):
        return "academic_thesis"
    if protected and not any(item["type"] in {"abstract", "references", "body"} for item in sections):
        return "academic_form"
    return "academic_template"


def _is_heading_like(paragraph: Any, text: str) -> bool:
    return str(getattr(paragraph.style, "name", "")).lower().startswith(("heading", "标题")) or len(text) <= 80


def _looks_like_cover(text: str, paragraphs: list[Any]) -> bool:
    compact = "".join(text.split())
    return _contains_field(text) or compact in {"课程论文", "毕业论文", "课程设计"} or any(_contains_field(p.text) for p in paragraphs)


def _contains_field(text: str) -> bool:
    lowered = text.lower()
    return any(marker.lower() in lowered for marker in _FIELD_MARKERS) or bool(re.search(r"[_＿]{2,}|：\s*$|:\s*$", text))


def _alignment_name(value: Any) -> str | None:
    if value is None:
        return None
    return getattr(value, "name", str(value).split(".")[-1]).lower()


def _clamp(value: float) -> float:
    return round(max(0.0, min(1.0, float(value))), 4)
