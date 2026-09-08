from __future__ import annotations

import re
from hashlib import sha1
from copy import deepcopy
from pathlib import Path
from typing import Any

from docx import Document

from .docx_formatter import replace_paragraph_text


AUTO_FIX = "AUTO_FIX"
SUGGEST_ONLY = "SUGGEST_ONLY"
HITL_REQUIRED = "HITL_REQUIRED"

_HIGH_RISK = re.compile(r"(?:\d+(?:\.\d+)?\s*[%％]|实验结果|实验数据|结论|引用|参考文献|文献|公式|研究方法|方法|定义|定理|模型|\b(?:Figure|Table|Fig\.?|Tab\.?)\s*\d+)", re.I)


def review_content(path: Path, ai_suggestions: list[dict[str, Any]] | None = None, *, source: str = "local") -> dict[str, Any]:
    document = Document(path)
    issues: list[dict[str, Any]] = []
    for index, paragraph in enumerate(document.paragraphs):
        text = paragraph.text
        if not text.strip():
            continue
        issues.extend(_deterministic_issues(index, text))
    for candidate in ai_suggestions or []:
        issue = _normalize_candidate(candidate, document, source="ai")
        if issue:
            issues.append(issue)
    issues = _deduplicate(issues)
    return _build_review(issues)


def apply_content_review(source: Path, output: Path, review: dict[str, Any]) -> dict[str, Any]:
    document = Document(source)
    changes: list[dict[str, Any]] = []
    for issue in review.get("issues", []):
        if issue.get("action_policy") != AUTO_FIX:
            continue
        index = issue.get("paragraph_index")
        original = str(issue.get("original_text") or "")
        suggested = str(issue.get("suggested_text") or "")
        if not isinstance(index, int) or not original or not suggested or index < 0 or index >= len(document.paragraphs):
            continue
        paragraph = document.paragraphs[index]
        if paragraph.text != original:
            continue
        before = paragraph.text
        after = suggested
        if after == before:
            continue
        replace_paragraph_text(paragraph, after)
        changes.append({
            "issue_id": issue.get("issue_id"), "paragraph_index": index, "locator": issue.get("locator"), "issue_type": issue["issue_type"], "before": before,
            "after": after, "reason": issue["reason"], "confidence": issue["confidence"],
            "source": issue["source"], "action": "auto_fix", "verification_status": "pending",
        })
    document.save(output)
    output_document = Document(output)
    verified = 0
    failed = 0
    for change in changes:
        index = change["paragraph_index"]
        actual = output_document.paragraphs[index].text if 0 <= index < len(output_document.paragraphs) else None
        expected = change["after"]
        change["verification_status"] = "verified" if actual == expected else "failed"
        change["verification_evidence"] = {"paragraph_index": index, "expected": expected, "actual": actual, "source": "re-read output DOCX"}
        verified += change["verification_status"] == "verified"
        failed += change["verification_status"] == "failed"
    suggestion_provenance = []
    for issue in review.get("issues", []):
        if issue.get("action_policy") == SUGGEST_ONLY:
            suggestion_provenance.append({"issue_id": issue.get("issue_id"), "paragraph_index": issue.get("paragraph_index"), "locator": issue.get("locator"), "before": issue.get("original_text"), "suggested": issue.get("suggested_text"), "reason": issue.get("reason"), "confidence": issue.get("confidence"), "source": issue.get("source"), "action": "suggestion", "status": "suggested", "verification_status": "original_unchanged"})
    hitl = [item for item in review.get("issues", []) if item.get("action_policy") == HITL_REQUIRED]
    review["provenance"] = {"auto_fixes": changes, "suggestions": suggestion_provenance, "accepted": [], "hitl": hitl}
    review["verification"] = {"total": len(changes), "verified": verified, "failed": failed, "suggestions_original_unchanged": True}
    review["decision"] = "HUMAN_REVIEW" if failed or hitl else "COMPLETE"
    review["decision_reason"] = "自动修正验证失败或存在高风险内容。" if failed or hitl else ("存在仅供参考的内容修改建议。" if suggestion_provenance else "内容规则检查通过。")
    return review


def apply_content_suggestion(source: Path, output: Path, issue: dict[str, Any]) -> dict[str, Any]:
    document = Document(source)
    index = issue.get("paragraph_index")
    original = str(issue.get("original_text") or issue.get("original") or "")
    suggested = str(issue.get("suggested_text") or issue.get("suggested") or "")
    issue_id = str(issue.get("issue_id") or "")
    if not isinstance(index, int) or index < 0 or index >= len(document.paragraphs) or not original or not suggested:
        return {"status": "error", "error": "invalid suggestion locator or text"}
    expected_id = _issue_id(index, issue.get("issue_type", "content"), original, suggested)
    if issue_id != expected_id:
        return {"status": "conflict", "error": "unknown or tampered suggestion issue_id"}
    policy = _make_issue(index, issue.get("issue_type", "content"), original, suggested, issue.get("reason", "用户确认的内容建议"), float(issue.get("confidence", 0.85)), issue.get("source", "user_confirmed"))["action_policy"]
    if policy != SUGGEST_ONLY:
        return {"status": "conflict", "error": "本地 risk policy 不允许采纳该建议"}
    paragraph = document.paragraphs[index]
    if paragraph.text != original:
        return {"status": "conflict", "error": "stale suggestion: 当前段落已变化", "issue_id": issue_id, "current": paragraph.text, "expected": original}
    before = paragraph.text
    replace_paragraph_text(paragraph, suggested)
    document.save(output)
    reread = Document(output)
    actual = reread.paragraphs[index].text if index < len(reread.paragraphs) else None
    status = "verified" if actual == suggested else "failed"
    change = {"issue_id": issue_id, "paragraph_index": index, "locator": {"kind": "paragraph", "paragraph_index": index}, "issue_type": issue.get("issue_type", "content"), "before": before, "after": actual, "suggested": suggested, "reason": issue.get("reason"), "confidence": issue.get("confidence"), "source": "user_confirmed", "action": "accepted_by_user", "status": "applied" if status == "verified" else "failed", "verification_status": status, "accepted_by_user": True, "verification_evidence": {"expected": suggested, "actual": actual, "source": "re-read output DOCX"}}
    return {"status": "ok" if status == "verified" else "failed", "output": output.name, "download_url": f"/download/{output.name}", "change": change}


def _deterministic_issues(index: int, text: str) -> list[dict[str, Any]]:
    found: list[dict[str, Any]] = []
    cleaned = text
    reasons = []
    if re.search(r"[ \t]{2,}", text):
        cleaned = re.sub(r"[ \t]{2,}", " ", cleaned)
        reasons.append("连续空格")
    if re.search(r"([。！？，；：])\1+", text):
        cleaned = re.sub(r"([。！？，；：])\1+", r"\1", cleaned)
        reasons.append("重复标点")
    residue = re.search(r"(?:\[?\s*[BC]-\d{1,3}\s*\]?|C-51)", text, re.I)
    if residue:
        cleaned = re.sub(r"(?:\[?\s*[BC]-\d{1,3}\s*\]?|C-51)", "", cleaned, flags=re.I)
        reasons.append("模板编号残留")
    if cleaned != text:
        found.append(_make_issue(index, "typo / punctuation" if len(reasons) < 3 else "formatting-like textual residue", text, cleaned, f"检测到{'、'.join(reasons)}，可安全规范化。", 0.99, "local"))
    for issue_type, original, suggested, reason, confidence in [("redundant expression", "很多的", "许多", "删去不必要的助词。", 0.92), ("overly colloquial wording", "大家都认为", "通常认为", "避免聊天式表述。", 0.9), ("awkward wording", "有很大的好处", "具有积极作用", "改为更正式的论文语体。", 0.88)]:
        if original in text:
            found.append(_make_issue(index, issue_type, text, text.replace(original, suggested, 1), reason, confidence, "local"))
    if _HIGH_RISK.search(text) and not any(item.get("action_policy") == HITL_REQUIRED for item in found):
        found.append(_make_issue(index, "high-risk content", text, text, "段落涉及数字、实验结果、结论、引用或方法定义，禁止自动改写。", 0.99, "local"))
    return found


def _normalize_candidate(item: dict[str, Any], document: Document, *, source: str) -> dict[str, Any] | None:
    try:
        index = int(item.get("paragraph_index", -1))
    except Exception:
        index = -1
    original_fragment = str(item.get("original", "")).strip()
    suggested = str(item.get("replacement", item.get("suggested", ""))).strip()
    if index < 0 or not original_fragment or not suggested or original_fragment == suggested or index >= len(document.paragraphs):
        return None
    original = document.paragraphs[index].text
    suggested_full = original.replace(original_fragment, suggested, 1)
    return _make_issue(index, str(item.get("issue_type", "unclear expression")), original, suggested_full, str(item.get("reason", "AI 发现可能需要人工确认的表达问题。")), float(item.get("confidence", 0.85) or 0.85), source)


def _make_issue(index: int, issue_type: str, original: str, suggested: str, reason: str, confidence: float, source: str) -> dict[str, Any]:
    high_risk = bool(_HIGH_RISK.search(original))
    large_change = len(suggested) > max(80, len(original) * 1.5) or len(original) > 80
    if issue_type in {"typo / punctuation", "formatting-like textual residue", "suspicious template residue"} and confidence >= 0.95 and not high_risk and not large_change:
        policy = AUTO_FIX
    elif high_risk or confidence < 0.8 or large_change:
        policy = HITL_REQUIRED
    else:
        policy = SUGGEST_ONLY
    locator = {"kind": "paragraph", "paragraph_index": index, "target_type": "body_paragraph"}
    return {"issue_id": _issue_id(index, issue_type, original, suggested), "paragraph_index": index, "locator": locator, "target_type": "body_paragraph", "issue_type": issue_type, "original_text": original, "evidence": reason, "reason": reason, "confidence": round(max(0.0, min(1.0, confidence)), 2), "risk_level": "high_risk" if policy == HITL_REQUIRED else "low" if policy == AUTO_FIX else "warning", "suggested_text": suggested, "action_policy": policy, "source": source, "status": "detected", "requires_user_confirmation": policy != AUTO_FIX}


def _issue_id(index: int, issue_type: Any, original: str, suggested: str) -> str:
    return f"con-{sha1(f'{index}|{issue_type}|{original}|{suggested}'.encode('utf-8')).hexdigest()[:12]}"


def _deduplicate(issues: list[dict[str, Any]]) -> list[dict[str, Any]]:
    result, seen = [], set()
    for issue in issues:
        key = (issue["paragraph_index"], issue["issue_type"], issue["original_text"], issue["suggested_text"])
        if key not in seen:
            seen.add(key); result.append(issue)
    return result[:80]


def _build_review(issues: list[dict[str, Any]]) -> dict[str, Any]:
    counts = {AUTO_FIX: 0, SUGGEST_ONLY: 0, HITL_REQUIRED: 0}
    for issue in issues:
        counts[issue["action_policy"]] += 1
    return {"issues": issues, "counts": counts, "auto_fix_count": counts[AUTO_FIX], "suggestion_count": counts[SUGGEST_ONLY], "hitl_count": counts[HITL_REQUIRED], "provenance": {"auto_fixes": [], "suggestions": [], "hitl": []}, "verification": {"total": 0, "verified": 0, "failed": 0, "suggestions_original_unchanged": True}}
