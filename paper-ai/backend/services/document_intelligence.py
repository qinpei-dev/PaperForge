from __future__ import annotations

import re
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any

from docx import Document

from .document_model import DocumentModel, extract_captions, extract_headings
from .docx_analyzer import count_inline_images


SECTION_TYPES = {
    "title",
    "abstract",
    "keywords",
    "introduction",
    "related_work",
    "methodology",
    "experiment",
    "conclusion",
    "references",
    "unknown",
}

SECTION_TERMS: dict[str, tuple[str, ...]] = {
    "abstract": ("摘要", "abstract", "summary"),
    "keywords": ("关键词", "关键字", "keywords", "key words"),
    "introduction": ("引言", "绪论", "绪言", "研究背景", "introduction", "background"),
    "related_work": ("相关工作", "文献综述", "研究现状", "相关研究", "related work", "literature review", "state of the art"),
    "methodology": ("研究方法", "方法", "方法论", "系统设计", "模型构建", "实验方法", "method", "methodology", "approach", "system design"),
    "experiment": ("实验", "实验结果", "实证分析", "结果与分析", "实验分析", "experiment", "experiments", "evaluation", "results", "analysis"),
    "conclusion": ("结论", "结语", "总结", "展望", "结论与展望", "conclusion", "conclusions", "future work"),
    "references": ("参考文献", "文献", "参考资料", "references", "bibliography", "works cited"),
}

_HEADING_NUMBER = re.compile(r"^(?:第\s*[一二三四五六七八九十百千万0-9]+\s*[章节篇部分]|[0-9]+(?:\.[0-9]+)*[.、：:]?)\s*")
_REFERENCE_LINE = re.compile(r"^(?:\[[0-9]+\]|[［【][0-9]+[］】]|[0-9]+[.、)]\s*)")
_EQUATION_LINE = re.compile(r"(?:=|∑|∫|≤|≥|\+\-|\^|\barg\s*min\b|\bwhere\b)")


@dataclass(frozen=True)
class DocumentSection:
    title: str
    type: str
    confidence: float
    paragraph_id: int | None = None
    level: int | None = None


@dataclass(frozen=True)
class DocumentElement:
    paragraph_id: int
    semantic_type: str
    confidence: float
    text: str = ""


@dataclass
class DocumentAnalysis:
    document_type: str
    sections: list[DocumentSection] = field(default_factory=list)
    elements: list[DocumentElement] = field(default_factory=list)
    confidence: float = 0.0
    metadata: dict[str, Any] = field(default_factory=dict)

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)


def analyze_document(
    source: Path | DocumentModel | dict[str, Any],
    *,
    document_path: Path | None = None,
) -> DocumentAnalysis:
    """Build a conservative semantic view without changing the existing model.

    A DocumentModel is preferred because it is already built by the live Agent.
    A path is also accepted for direct use by tests and future callers.
    """
    if isinstance(source, DocumentModel):
        model = source
        document = None
    elif isinstance(source, (str, Path)):
        document = Document(source)
        paragraphs = list(document.paragraphs)
        nonempty = [(i, p) for i, p in enumerate(paragraphs) if p.text.strip()]
        headings, _ = extract_headings(nonempty)
        captions = extract_captions(nonempty)
        model = None
    elif isinstance(source, dict):
        return _analysis_from_dict(source)
    else:
        raise TypeError("source must be a Path, DocumentModel, or mapping")

    if model is not None:
        # The model intentionally contains only serializable structure.  Reopen
        # the document for paragraph text and tables while reusing its heading
        # and caption decisions.
        document = Document(document_path) if document_path is not None else None
        if document is None:
            raise ValueError("DocumentModel analysis requires document_path")
        paragraphs = list(document.paragraphs)
        headings = list(model.headings)
        captions = list(model.captions)
    assert document is not None

    heading_by_id = {item.paragraph_index: item for item in headings}
    caption_by_id = {item.paragraph_index: item for item in captions}
    reference_start = _find_reference_start(paragraphs)
    sections: list[DocumentSection] = []
    elements: list[DocumentElement] = []

    for index, paragraph in enumerate(paragraphs):
        text = paragraph.text.strip()
        if not text:
            continue
        heading = heading_by_id.get(index)
        section_type, section_confidence = classify_section_heading(text, heading)
        if heading is not None:
            title_type = "title" if _looks_like_title(text, index, paragraphs, heading) else section_type
            sections.append(DocumentSection(text, title_type, section_confidence, index, heading.level))
            elements.append(DocumentElement(index, "heading", heading.confidence, text))
            continue
        if index < reference_start:
            semantic_type, confidence = _classify_non_heading(text, index, paragraphs, caption_by_id)
        else:
            semantic_type, confidence = "reference", 0.94 if _REFERENCE_LINE.search(text) else 0.78
        elements.append(DocumentElement(index, semantic_type, confidence, text))

    # Title/abstract/keywords are sometimes plain paragraphs, not heading styles.
    sections = _add_labeled_sections(sections, paragraphs)
    if paragraphs:
        first_index, first_paragraph = next(((i, p) for i, p in enumerate(paragraphs) if p.text.strip()), (None, None))
        if first_index is not None and not any(item.type == "title" for item in sections):
            first_text = first_paragraph.text.strip()
            if len(first_text) <= 180 and classify_section_heading(first_text)[0] == "unknown":
                sections.insert(0, DocumentSection(first_text, "title", 0.68, first_index, None))
    known = [item.confidence for item in sections + elements]
    overall = _clamp(sum(known) / len(known) if known else 0.0)
    metadata = {
        "paragraph_count": len(paragraphs),
        "nonempty_paragraph_count": sum(bool(p.text.strip()) for p in paragraphs),
        "heading_count": len(headings),
        "table_count": len(document.tables),
        "figure_count": model.figure_image_count if model is not None else count_inline_images(document),
        "reference_count": max(0, sum(1 for i, p in enumerate(paragraphs) if i >= reference_start and p.text.strip())),
        "reference_start_paragraph": reference_start if reference_start < len(paragraphs) else None,
        "language": _detect_language(" ".join(p.text for p in paragraphs if p.text.strip())),
        "warnings": [] if sections else ["未可靠识别论文章节结构。"],
    }
    return DocumentAnalysis(model.document_type if model is not None else "unknown", sections, elements, overall, metadata)


def classify_section_heading(text: str, heading: Any | None = None) -> tuple[str, float]:
    normalized = _normalize_heading(text)
    if not normalized:
        return "unknown", 0.25
    candidates = [(kind, term) for kind, terms in SECTION_TERMS.items() for term in terms if _term_matches(normalized, term)]
    if candidates:
        kind, term = max(candidates, key=lambda item: len(item[1]))
        base = 0.93 if heading is not None else 0.78
        if _HEADING_NUMBER.match(text):
            base += 0.03
        return kind, _clamp(base)
    return "unknown", 0.52 if heading is not None else 0.30


def _classify_non_heading(text: str, index: int, paragraphs: list[Any], captions: dict[int, Any]) -> tuple[str, float]:
    if index in captions:
        return "caption", captions[index].confidence
    lowered = text.lower()
    if _EQUATION_LINE.search(text) and len(text) < 240:
        return "equation", 0.70
    if lowered.startswith(("图", "图 ", "图：", "figure", "fig.")) and len(text) < 180:
        return "caption", 0.84
    if lowered.startswith(("表", "表 ", "表：", "table")) and len(text) < 180:
        return "table_description", 0.84
    if _REFERENCE_LINE.search(text):
        return "reference", 0.82
    return "body", 0.90 if len(text) >= 20 else 0.66


def _add_labeled_sections(sections: list[DocumentSection], paragraphs: list[Any]) -> list[DocumentSection]:
    found = {(item.type, item.paragraph_id) for item in sections}
    result = list(sections)
    for index, paragraph in enumerate(paragraphs):
        text = paragraph.text.strip()
        if not text:
            continue
        kind, confidence = classify_section_heading(text)
        if kind in {"abstract", "keywords"} and (kind, index) not in found:
            result.append(DocumentSection(text, kind, confidence, index, None))
    return sorted(result, key=lambda item: item.paragraph_id if item.paragraph_id is not None else -1)


def _find_reference_start(paragraphs: list[Any]) -> int:
    for index, paragraph in enumerate(paragraphs):
        kind, confidence = classify_section_heading(paragraph.text.strip())
        if kind == "references" and confidence >= 0.5:
            return index + 1
    return len(paragraphs)


def _looks_like_title(text: str, index: int, paragraphs: list[Any], heading: Any) -> bool:
    if index != 0 or (heading.level is not None and heading.level > 1):
        return False
    return len(text) <= 120 and not any(_term_matches(_normalize_heading(text), term) for terms in SECTION_TERMS.values() for term in terms)


def _normalize_heading(text: str) -> str:
    value = _HEADING_NUMBER.sub("", text.strip()).strip(" ：:.-")
    return re.sub(r"\s+", " ", value).lower()


def _term_matches(text: str, term: str) -> bool:
    normalized_term = term.lower()
    return text == normalized_term or text.startswith(normalized_term + " ") or text.startswith(normalized_term + "：") or text.startswith(normalized_term + ":")


def _detect_language(text: str) -> str:
    return "zh" if len(re.findall(r"[\u4e00-\u9fff]", text)) >= len(re.findall(r"[A-Za-z]", text)) else "en"


def _clamp(value: float) -> float:
    return round(max(0.0, min(1.0, value)), 4)


def _analysis_from_dict(value: dict[str, Any]) -> DocumentAnalysis:
    sections = [DocumentSection(**item) for item in value.get("sections", []) if isinstance(item, dict)]
    elements = [DocumentElement(**item) for item in value.get("elements", []) if isinstance(item, dict)]
    return DocumentAnalysis(str(value.get("document_type") or "unknown"), sections, elements, float(value.get("confidence") or 0.0), dict(value.get("metadata") or {}))
