from __future__ import annotations

from dataclasses import asdict, dataclass, field
from hashlib import sha256
from pathlib import Path
from typing import Any

from docx import Document

from .docx_analyzer import (
    analyze_references,
    count_inline_images,
    get_body_paragraph_indices,
    guess_heading_level,
    is_section_heading,
    parse_caption_number,
)


@dataclass(frozen=True)
class HeadingNode:
    level: int | None
    text: str
    paragraph_index: int
    style: str | None
    confidence: float


@dataclass(frozen=True)
class CaptionNode:
    kind: str
    text: str
    paragraph_index: int
    number: str | int
    confidence: float


@dataclass
class DocumentModel:
    document_id: str
    document_type: str
    paragraph_count: int
    section_count: int
    table_count: int
    figure_image_count: int
    headings: list[HeadingNode]
    captions: list[CaptionNode]
    abstract: str | None
    keywords: list[str]
    references: list[str]
    structural_fingerprint: dict[str, Any]
    style_summary: dict[str, Any]
    body_paragraph_indices: list[int] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)


def build_document_model(
    path: Path,
    *,
    classification: dict[str, Any] | None = None,
) -> DocumentModel:
    """Build a conservative, serializable DOCX representation for planning.

    This deliberately exposes only facts read from the document or existing
    classifier output. Ambiguous structure is represented by confidence and
    warnings instead of being promoted to a certain classification.
    """
    document = Document(path)
    paragraphs = list(document.paragraphs)
    nonempty = [(index, paragraph) for index, paragraph in enumerate(paragraphs) if paragraph.text.strip()]
    headings, heading_warnings = extract_headings(nonempty)
    captions = extract_captions(nonempty)
    abstract, abstract_warning = extract_abstract(nonempty)
    keywords, keyword_warning = extract_keywords(nonempty)
    reference_check = analyze_references([paragraph.text.strip() for _, paragraph in nonempty])
    references = extract_reference_text(nonempty)
    warnings = [*heading_warnings]
    if abstract_warning:
        warnings.append(abstract_warning)
    if keyword_warning:
        warnings.append(keyword_warning)
    if not reference_check.get("has_reference_section"):
        warnings.append("未可靠识别参考文献章节。")

    style_counts: dict[str, int] = {}
    for _, paragraph in nonempty:
        style_name = paragraph.style.name if paragraph.style else "unknown"
        style_counts[style_name] = style_counts.get(style_name, 0) + 1

    fingerprint = build_structural_fingerprint(
        paragraphs=paragraphs,
        headings=headings,
        table_count=len(document.tables),
        image_count=count_inline_images(document),
        section_count=len(document.sections),
    )
    document_type = str((classification or {}).get("document_type") or "unknown")
    return DocumentModel(
        document_id=file_identifier(path),
        document_type=document_type,
        paragraph_count=len(paragraphs),
        section_count=len(document.sections),
        table_count=len(document.tables),
        figure_image_count=count_inline_images(document),
        headings=headings,
        captions=captions,
        abstract=abstract,
        keywords=keywords,
        references=references,
        structural_fingerprint=fingerprint,
        style_summary={
            "nonempty_paragraph_count": len(nonempty),
            "paragraph_style_counts": style_counts,
            "normal_style": read_normal_style(document),
        },
        body_paragraph_indices=get_body_paragraph_indices(document),
        warnings=list(dict.fromkeys(warnings)),
    )


def extract_headings(paragraphs: list[tuple[int, Any]]) -> tuple[list[HeadingNode], list[str]]:
    nodes: list[HeadingNode] = []
    for index, paragraph in paragraphs:
        text = paragraph.text.strip()
        style_name = paragraph.style.name if paragraph.style else None
        style_level = heading_level_from_style(style_name)
        heuristic_heading = is_section_heading(text)
        if not style_level and not heuristic_heading:
            continue
        level = style_level or guess_heading_level(text)
        confidence = 0.95 if style_level else 0.72
        nodes.append(HeadingNode(level=level, text=text, paragraph_index=index, style=style_name, confidence=confidence))
    warnings = [] if nodes else ["未可靠识别章节标题；标题层级不能作为自动修改依据。"]
    return nodes, warnings


def extract_captions(paragraphs: list[tuple[int, Any]]) -> list[CaptionNode]:
    """Keep only captions whose leading number is parsed by the existing analyzer."""
    captions: list[CaptionNode] = []
    for index, paragraph in paragraphs:
        text = paragraph.text.strip()
        for kind in ("figure", "table"):
            number = parse_caption_number(text, kind)
            if number is not None:
                captions.append(CaptionNode(kind=kind, text=text, paragraph_index=index, number=number, confidence=0.95))
                break
    return captions


def heading_level_from_style(style_name: str | None) -> int | None:
    if not style_name:
        return None
    normalized = style_name.lower().replace(" ", "")
    for level in (1, 2, 3):
        if normalized in {f"heading{level}", f"标题{level}"}:
            return level
    return None


def extract_abstract(paragraphs: list[tuple[int, Any]]) -> tuple[str | None, str | None]:
    for position, (_, paragraph) in enumerate(paragraphs):
        text = paragraph.text.strip()
        if text.startswith(("摘要：", "摘要:", "Abstract:")):
            return text.split(":", 1)[-1].split("：", 1)[-1].strip() or None, None
        if text.lower() in {"摘要", "abstract"} and position + 1 < len(paragraphs):
            candidate = paragraphs[position + 1][1].text.strip()
            return candidate or None, None
    return None, "未可靠识别摘要。"


def extract_keywords(paragraphs: list[tuple[int, Any]]) -> tuple[list[str], str | None]:
    for _, paragraph in paragraphs:
        text = paragraph.text.strip()
        if text.startswith(("关键词：", "关键词:", "关键字：", "Keywords:", "keywords:")):
            raw = text.split(":", 1)[-1].split("：", 1)[-1]
            values = [item.strip() for item in raw.replace("；", ";").replace("、", ";").split(";") if item.strip()]
            return values, None
    return [], "未可靠识别关键词。"


def extract_reference_text(paragraphs: list[tuple[int, Any]]) -> list[str]:
    in_references = False
    values: list[str] = []
    for _, paragraph in paragraphs:
        text = paragraph.text.strip()
        if text.rstrip("：:").lower() in {"参考文献", "references", "bibliography"}:
            in_references = True
            continue
        if in_references and text:
            values.append(text)
    return values


def read_normal_style(document: Document) -> dict[str, Any]:
    try:
        style = document.styles["Normal"]
        return {"font_name": style.font.name, "font_size_pt": style.font.size.pt if style.font.size else None}
    except Exception:
        return {"font_name": None, "font_size_pt": None, "warning": "Normal 样式不可读"}


def build_structural_fingerprint(
    *,
    paragraphs: list[Any],
    headings: list[HeadingNode],
    table_count: int,
    image_count: int,
    section_count: int,
) -> dict[str, Any]:
    payload = {
        "paragraph_count": len(paragraphs),
        "nonempty_paragraph_count": sum(1 for paragraph in paragraphs if paragraph.text.strip()),
        "heading_outline": [(node.level, node.text) for node in headings],
        "table_count": table_count,
        "image_count": image_count,
        "section_count": section_count,
    }
    digest = sha256(repr(payload).encode("utf-8")).hexdigest()[:16]
    return {**payload, "signature": digest}


def file_identifier(path: Path) -> str:
    digest = sha256(path.read_bytes()).hexdigest()[:16]
    return f"docx-{digest}"
