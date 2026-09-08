from __future__ import annotations

from pathlib import Path
from tempfile import TemporaryDirectory

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

from services.document_model import build_document_model
from services.docx_analyzer import analyze_docx
from services.executor_adapter import execute_plan
from services.planner import ExecutionPlan, PlanStep
from services.rule_engine import Rule, RuleSource
from services.verifier import verify_output


def check(name: str, condition: bool) -> None:
    if not condition:
        raise AssertionError(f"{name} FAIL")
    print(f"{name} PASS")


def make_source(path: Path) -> None:
    document = Document()
    document.add_paragraph("正文目标级验证")
    first = document.add_paragraph("第一段目标正文。")
    first.alignment = WD_ALIGN_PARAGRAPH.LEFT
    first.paragraph_format.first_line_indent = Cm(0)
    first.runs[0].font.size = Pt(9)
    second = document.add_paragraph("第二段非目标正文。")
    second.alignment = WD_ALIGN_PARAGRAPH.LEFT
    second.runs[0].font.size = Pt(9)
    document.add_paragraph("参考文献")
    document.add_paragraph("[1] 示例[J]. 测试, 2026.")
    document.save(path)


def plan_step(step_id: str, rule: Rule, indices: list[int] | None) -> PlanStep:
    locator = {"kind": "paragraph_indices", "indices": indices, "semantic_role": "body", "target_type": "body_paragraph"} if indices is not None else {}
    return PlanStep(step_id, "apply_body_font" if rule.property in {"font_name", "font_size"} else "apply_body_paragraph_format", "body", rule.id, "test", "low", True, [], "planned", "test", target_locator=locator)


def main() -> None:
    with TemporaryDirectory() as directory:
        source, output = Path(directory) / "source.docx", Path(directory) / "output.docx"
        make_source(source)
        model = build_document_model(source, classification={"document_type": "academic_paper"})
        before = analyze_docx(source)
        target_index = 1
        rules = [
            Rule("body-size", "body", "font_size", 12, RuleSource.DEFAULT, "test", .9, "low", True),
            Rule("body-alignment", "body", "alignment", WD_ALIGN_PARAGRAPH.JUSTIFY, RuleSource.DEFAULT, "test", .9, "low", True),
            Rule("body-indent", "body", "first_line_indent_cm", .74, RuleSource.DEFAULT, "test", .9, "low", True),
            Rule("body-space-after", "body", "space_after_pt", 6, RuleSource.DEFAULT, "test", .9, "low", True),
        ]
        plan = ExecutionPlan("plan-p1-2", model.document_id, [plan_step(f"step-{rule.id}", rule, [target_index]) for rule in rules], [], False)
        execution = execute_plan(plan, source, output, None, rules)
        check("body_targets_executed", len(execution.changes) == len(rules) and all(change["target"]["paragraph_index"] == target_index and change["target"]["target_type"] == "body_paragraph" for change in execution.changes))
        output_doc = Document(output)
        check("only_target_paragraph_changed", output_doc.paragraphs[2].alignment == WD_ALIGN_PARAGRAPH.LEFT and output_doc.paragraphs[2].runs[0].font.size.pt == 9)
        verification = verify_output(model, before, output, None, rules, plan, execution.unsupported_step_ids, execution.changes)
        check("body_targets_reloaded_and_verified", all(change["verification_scope"] == "target" and change["verification_status"] == "verified" and change["verification_evidence"]["paragraph_index"] == target_index and change["verification_evidence"]["source"] == "re-read output DOCX" for change in verification.provenance_changes))

        mismatched_changes = [dict(change, expected=999) for change in execution.changes]
        mismatch = verify_output(model, before, output, None, rules, plan, execution.unsupported_step_ids, mismatched_changes)
        check("expected_actual_mismatch_not_verified", all(change["verification_status"] == "failed" for change in mismatch.provenance_changes))

        unsupported_plan = ExecutionPlan("plan-p1-2-unsupported", model.document_id, [plan_step("missing-body-target", rules[0], None)], [], False)
        unsupported_output = Path(directory) / "unsupported.docx"
        unsupported_execution = execute_plan(unsupported_plan, source, unsupported_output, None, rules)
        unsupported_verification = verify_output(model, before, unsupported_output, None, rules, unsupported_plan, unsupported_execution.unsupported_step_ids, unsupported_execution.changes)
        check("unsupported_body_target_honest", unsupported_execution.step_statuses["missing-body-target"] == "unsupported" and unsupported_verification.provenance_changes[0]["verification_status"] == "unsupported" and "reason" in unsupported_verification.provenance_changes[0]["verification_evidence"])
    print("P1_2_BODY_TARGET_VERIFICATION PASS")


if __name__ == "__main__":
    main()
