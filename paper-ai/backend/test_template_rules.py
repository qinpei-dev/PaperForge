from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from services.ai_reasoning import generate_reasoning
from services.template_intelligence import analyze_template, build_effective_rules


def _make_rule_template(path: Path) -> None:
    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    inherited = document.styles.add_style("Inherited Paragraph", WD_STYLE_TYPE.PARAGRAPH)
    inherited.base_style = normal
    document.add_paragraph("Body using an inherited paragraph style.", style="Inherited Paragraph")

    cover = document.add_paragraph("毕业论文")
    cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cover.runs[0].font.name = "SimHei"
    cover.runs[0].font.size = Pt(18)
    document.add_paragraph("姓名：________")

    for title, font_name, size in (("Abstract", "Calibri", 14), ("References", "Times New Roman", 10), ("Appendix", "Courier New", 12)):
        paragraph = document.add_paragraph(title, style="Heading 1")
        paragraph.runs[0].font.name = font_name
        paragraph.runs[0].font.size = Pt(size)
        paragraph.runs[0].font.bold = True
    document.save(path)


def test_style_inheritance_and_region_overrides(tmp_path: Path) -> None:
    path = tmp_path / "rules.docx"
    _make_rule_template(path)
    analysis = analyze_template(path)

    inherited = [item for item in analysis.inherited_rules if item["scope"] == "style:Inherited Paragraph" and item["property"] == "font"]
    assert inherited and inherited[0]["value"] == "Arial"
    assert inherited[0]["inherits_from"] == "Normal"
    assert {item["scope"] for item in analysis.override_rules} >= {"abstract", "references", "appendix", "cover"}
    assert analysis.effective_rules["abstract"]["font"] == "Calibri"
    assert analysis.effective_rules["references"]["size"] == 10
    assert analysis.effective_rules["appendix"]["font"] == "Courier New"
    assert analysis.effective_rules["cover"]["protected"] is True


def test_effective_rule_priority_is_deterministic() -> None:
    effective = build_effective_rules(
        [{"scope": "body", "property": "font", "value": "Base", "priority": 10}],
        [{"scope": "body", "property": "font", "value": "Inherited", "priority": 20}],
        [{"scope": "body", "property": "font", "value": "Override", "priority": 30}],
        {"template_base": 10, "style_inheritance": 20, "section_override": 30},
    )
    assert effective["body"]["font"] == "Override"


def test_reasoning_prefers_effective_rules_and_keeps_legacy_call_compatible() -> None:
    legacy = generate_reasoning({"confidence": 0.8}, [{"issue_type": "paragraph_format_mismatch", "score": 70}], [])
    result = generate_reasoning(
        {"confidence": 0.8},
        [{"issue_type": "paragraph_format_mismatch", "score": 70}],
        [],
        {"confidence": 0.9, "effective_rules": {"body": {"font": "Arial", "size": 11}}},
    )
    assert legacy
    assert result[0]["context"]["effective_rule_scope"] == "body"
    assert result[0]["context"]["effective_rule"] == {"font": "Arial", "size": 11}

