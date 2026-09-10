from __future__ import annotations

import io
import time
from collections.abc import Generator
from pathlib import Path

import pytest
from docx import Document
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import Session, sessionmaker

import main
from db.base import Base
from db.session import get_db
from services.storage import LocalTemplateStorage
from services.task_worker import TaskWorker
from services.template_persistence import refresh_template_registry
from services.template_registry import BUILTIN_TEMPLATE_DEFINITIONS, template_registry


MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def docx_bytes() -> bytes:
    document = Document()
    document.add_heading("模板标题", 0)
    document.add_paragraph("这是可供 Template Intelligence 解析的正文。")
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


@pytest.fixture(autouse=True)
def reset_registry() -> Generator[None, None, None]:
    template_registry.replace(BUILTIN_TEMPLATE_DEFINITIONS)
    yield
    template_registry.replace(BUILTIN_TEMPLATE_DEFINITIONS)


@pytest.fixture
def managed_client(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> Generator[tuple[TestClient, sessionmaker[Session], LocalTemplateStorage], None, None]:
    monkeypatch.setenv("AUTH_REQUIRED", "true")
    monkeypatch.setenv("AUTO_CREATE_DB", "false")
    engine = create_engine(f"sqlite:///{(tmp_path / 'managed.db').as_posix()}")
    Base.metadata.create_all(engine)
    sessions = sessionmaker(bind=engine, autoflush=False, expire_on_commit=False)
    storage = LocalTemplateStorage(tmp_path / "template-storage")
    monkeypatch.setattr(main, "MANAGED_TEMPLATE_STORAGE", storage)
    monkeypatch.setattr("services.template_persistence._TEMPLATE_STORAGE", storage)

    def override_get_db() -> Generator[Session, None, None]:
        db = sessions()
        try:
            yield db
        finally:
            db.close()

    app = main.app
    app.dependency_overrides[get_db] = override_get_db
    monkeypatch.setattr(main, "SessionLocal", sessions)
    worker = TaskWorker()
    monkeypatch.setattr(main, "task_worker", worker)
    with TestClient(app) as client:
        try:
            yield client, sessions, storage
        finally:
            worker.shutdown()
    app.dependency_overrides.clear()
    engine.dispose()


def account(client: TestClient, email: str) -> dict[str, str]:
    response = client.post("/auth/register", json={"email": email, "password": "password123"})
    assert response.status_code == 201, response.text
    return {"Authorization": f"Bearer {response.json()['access_token']}"}


def upload(client: TestClient, headers: dict[str, str], *, template_id: str = "tenant-thesis", version: str = "1.0") -> dict[str, object]:
    response = client.post("/templates", headers=headers, data={"name": "Tenant Thesis", "template_id": template_id, "version": version, "school": "Test U", "document_type": "academic_paper"}, files={"file": ("thesis.docx", docx_bytes(), MIME)})
    assert response.status_code == 201, response.text
    return response.json()


def test_storage_rejects_traversal_and_is_tenant_isolated(tmp_path: Path) -> None:
    storage = LocalTemplateStorage(tmp_path / "storage")
    first = storage.save("tenant-a", "resource-a", "1.0", io.BytesIO(b"a"))
    second = storage.save("tenant-b", "resource-b", "1.0", io.BytesIO(b"b"))
    assert first != second and storage.open(first).read() == b"a" and storage.open(second).read() == b"b"
    with pytest.raises(ValueError):
        storage.resolve_local_path("local://tenant-templates/../escape/1.0/file.docx")
    with pytest.raises(ValueError):
        storage.resolve_local_path("local://tenant-templates/%2e%2e/escape/1.0/file.docx")
    storage.delete(first)
    assert not storage.exists(first) and storage.exists(second)


def test_upload_lifecycle_and_tenant_isolation(managed_client) -> None:
    client, _, storage = managed_client
    a, b = account(client, "p3-a@example.com"), account(client, "p3-b@example.com")
    created = upload(client, a)
    resource_id = str(created["id"])
    assert created["scope"] == "tenant" and created["tenant_id"]
    assert any(item["id"] == resource_id for item in client.get("/templates", headers=a).json()["templates"])
    assert all(item["id"] != resource_id for item in client.get("/templates", headers=b).json()["templates"])
    assert client.get(f"/templates/{resource_id}", headers=b).status_code == 404
    assert client.get(f"/templates/{resource_id}/file", headers=b).status_code == 404
    assert client.patch(f"/templates/{resource_id}", headers=b, json={"name": "no"}).status_code == 404
    assert client.delete(f"/templates/{resource_id}", headers=b).status_code == 404
    assert client.get(f"/templates/{resource_id}/file", headers=a).content.startswith(b"PK")
    changed = client.patch(f"/templates/{resource_id}", headers=a, json={"name": "Renamed", "status": "disabled"})
    assert changed.status_code == 200 and changed.json()["template_id"] == "tenant-thesis"
    disabled = client.post("/tasks", headers=a, data={"mode": "local", "allow_non_paper": "true", "template_id": "tenant-thesis", "template_version": "1.0"}, files={"paper": ("paper.docx", docx_bytes(), MIME)})
    assert disabled.status_code == 422
    assert upload(client, a, version="1.1")["version"] == "1.1"
    assert client.post("/templates", headers=a, data={"name": "duplicate", "template_id": "tenant-thesis", "version": "1.1"}, files={"file": ("same.docx", docx_bytes(), MIME)}).status_code == 409
    assert storage.root.exists()


def test_invalid_upload_and_task_execution_resolve_storage_locator(managed_client, monkeypatch: pytest.MonkeyPatch) -> None:
    client, sessions, _ = managed_client
    headers = account(client, "p3-run@example.com")
    invalid = client.post("/templates", headers=headers, data={"name": "bad", "version": "1"}, files={"file": ("bad.pdf", b"no", "application/pdf")})
    assert invalid.status_code == 400
    corrupt = client.post("/templates", headers=headers, data={"name": "bad", "version": "1"}, files={"file": ("bad.docx", b"no", MIME)})
    assert corrupt.status_code == 422
    created = upload(client, headers, template_id="run-template")
    captured: list[Path | None] = []

    def fake_pipeline(**kwargs: object) -> dict[str, object]:
        selected = kwargs["resolved_template"]
        captured.append(selected.template_path)
        output = Path(kwargs["output_dir"]) / "p3-result.docx"
        output.write_bytes(docx_bytes())
        return {"status": "ok", "filename": output.name, "after_score": 90.0, "agent_trace": [], "modification_report": {"summary": "ok"}, "template": selected.provenance()}

    monkeypatch.setattr(main, "run_agent_pipeline", fake_pipeline)
    response = client.post("/tasks", headers=headers, data={"mode": "local", "allow_non_paper": "true", "template_id": "run-template", "template_version": "1.0"}, files={"paper": ("paper.docx", docx_bytes(), MIME)})
    assert response.status_code == 201
    task_id = response.json()["task_id"]
    for _ in range(50):
        detail = client.get(f"/tasks/{task_id}", headers=headers).json()
        if detail["status"] == "completed":
            break
        time.sleep(0.02)
    assert captured and captured[0] is not None and captured[0].is_file()
    assert detail["template"]["id"] == "run-template"
    assert client.delete(f"/templates/{created['id']}", headers=headers).status_code == 409
    db = sessions()
    refresh_template_registry(db)
    db.close()
