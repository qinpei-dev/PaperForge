from pathlib import Path
from tempfile import TemporaryDirectory

from docx import Document

from services.content_review import AUTO_FIX, HITL_REQUIRED, SUGGEST_ONLY, apply_content_review, review_content


def make_doc(path: Path, paragraphs: list[str]) -> None:
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    doc.save(path)


def main() -> None:
    with TemporaryDirectory() as tmp:
        root = Path(tmp)
        source = root / "source.docx"
        make_doc(source, ["这是  一个测试。。", "很多的表达需要调整。", "实验结果为 95%，结论需要人工确认。", "其它段落保持不变。"])
        review = review_content(source)
        assert review["counts"][AUTO_FIX] == 1, review
        assert review["counts"][SUGGEST_ONLY] == 1, review
        assert review["counts"][HITL_REQUIRED] == 1, review
        output = root / "output.docx"
        result = apply_content_review(source, output, review)
        paragraphs = [p.text for p in Document(output).paragraphs]
        assert paragraphs[0] == "这是 一个测试。", paragraphs
        assert paragraphs[1] == "很多的表达需要调整。", paragraphs
        assert paragraphs[2] == "实验结果为 95%，结论需要人工确认。", paragraphs
        assert result["provenance"]["auto_fixes"][0]["verification_status"] == "verified", result
        assert result["provenance"]["suggestions"][0]["verification_status"] == "original_unchanged", result
        assert result["decision"] == "HUMAN_REVIEW", result

        ai_review = review_content(source, [{"paragraph_index": 1, "original": "很多的", "replacement": "大量的"}], source="ai")
        ai_issue = next(item for item in ai_review["issues"] if item["source"] == "ai")
        assert ai_issue["action_policy"] == SUGGEST_ONLY, ai_review
        ai_review = review_content(source, [{"paragraph_index": 2, "original": "实验结果为 95%", "replacement": "实验结果显著优于预期，证明方法有效"}], source="ai")
        ai_issue = next(item for item in ai_review["issues"] if item["source"] == "ai")
        assert ai_issue["action_policy"] == HITL_REQUIRED, ai_review
    print("P2_CONTENT_REVIEW_CLOSURE PASS")


if __name__ == "__main__":
    main()
